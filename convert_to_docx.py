#!/usr/bin/env python3
"""
Converts /home/user/hello-world/hausarbeit_final.md to
/home/user/hello-world/Hausarbeit_Softwaremodellierung_final.docx

Requires: python-docx  (pip install python-docx)
Run with: python3 /home/user/hello-world/convert_to_docx.py
"""

import re
import zipfile
import io
import os

# Try python-docx first; fall back to raw XML/ZIP generation
try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


# ═══════════════════════════════════════════════════════════════════════════
# PATH CONSTANTS
# ═══════════════════════════════════════════════════════════════════════════
MD_PATH   = '/home/user/hello-world/hausarbeit_final.md'
DOCX_PATH = '/home/user/hello-world/Hausarbeit_Softwaremodellierung_final.docx'


# ═══════════════════════════════════════════════════════════════════════════
# PYTHON-DOCX PATH  (preferred)
# ═══════════════════════════════════════════════════════════════════════════

def set_page_margins(doc):
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)


def configure_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after  = Pt(6)
    normal.paragraph_format.line_spacing = Pt(14)

    for lvl, sz in {1: 16, 2: 14, 3: 12, 4: 11}.items():
        st = doc.styles[f'Heading {lvl}']
        st.font.name  = 'Calibri'
        st.font.size  = Pt(sz)
        st.font.bold  = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before    = Pt(14 if lvl <= 2 else 8)
        st.paragraph_format.space_after     = Pt(4)
        st.paragraph_format.keep_with_next  = True


def set_table_borders(table):
    tbl   = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    borders = OxmlElement('w:tblBorders')
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{name}')
        el.set(qn('w:val'),   'single')
        el.set(qn('w:sz'),    '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tblPr.append(borders)


def shade_cell(cell, fill='D9D9D9'):
    tcPr = cell._tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  fill)
    tcPr.append(shd)


INLINE_RE = re.compile(r'(\*\*(.+?)\*\*|\*(.+?)\*|([^*]+))', re.DOTALL)


def add_inline(para, text, size=11):
    for m in INLINE_RE.finditer(text):
        bt = m.group(2)   # bold
        it = m.group(3)   # italic
        pt = m.group(4)   # plain
        run = para.add_run(bt or it or pt or '')
        run.font.name = 'Calibri'
        run.font.size = Pt(size)
        if bt: run.bold   = True
        if it: run.italic = True


def parse_table(lines, start):
    rows, i = [], start
    while i < len(lines):
        s = lines[i].strip()
        if not s or not s.startswith('|'):
            break
        if re.match(r'^\|[-|: ]+\|$', s):
            i += 1
            continue
        rows.append([c.strip() for c in s.strip('|').split('|')])
        i += 1
    return rows, i


def build_with_docx(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()
    set_page_margins(doc)
    configure_styles(doc)

    i = 0
    while i < len(lines):
        s = lines[i].strip()

        if not s:
            i += 1
            continue

        # horizontal rule
        if re.match(r'^-{3,}$', s):
            p = doc.add_paragraph()
            run = p.add_run('─' * 72)
            run.font.name  = 'Calibri'
            run.font.size  = Pt(7)
            run.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after  = Pt(8)
            i += 1
            continue

        # heading
        hm = re.match(r'^(#{1,4})\s+(.*)', s)
        if hm:
            lvl  = len(hm.group(1))
            text = hm.group(2).strip()
            heading = doc.add_heading('', level=lvl)
            heading.clear()
            fs = {1: 16, 2: 14, 3: 12, 4: 11}[lvl]
            add_inline(heading, text, size=fs)
            for run in heading.runs:
                run.bold = True
            i += 1
            continue

        # table
        if s.startswith('|'):
            rows, i = parse_table(lines, i)
            if not rows:
                continue
            ncols = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=ncols)
            tbl.style = 'Table Grid'
            set_table_borders(tbl)
            for ri, row in enumerate(rows):
                for ci in range(ncols):
                    cell = tbl.cell(ri, ci)
                    txt  = row[ci] if ci < len(row) else ''
                    p    = cell.paragraphs[0]
                    p.clear()
                    add_inline(p, txt, size=10)
                    if ri == 0:
                        shade_cell(cell)
                        for run in p.runs:
                            run.bold = True
            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(4)
            continue

        # normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(6)
        p.paragraph_format.line_spacing = Pt(14)
        add_inline(p, s, size=11)
        i += 1

    doc.save(docx_path)


# ═══════════════════════════════════════════════════════════════════════════
# FALLBACK: raw Office Open XML / ZIP  (no external deps)
# ═══════════════════════════════════════════════════════════════════════════

def esc(text):
    """XML-escape a string."""
    return (text
            .replace('&', '&amp;')
            .replace('<', '&lt;')
            .replace('>', '&gt;')
            .replace('"', '&quot;'))


def make_run(text, bold=False, italic=False, size_half=22):
    """Return a <w:r> XML fragment."""
    rpr = f'<w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="{size_half}"/><w:szCs w:val="{size_half}"/>'
    if bold:   rpr += '<w:b/>'
    if italic: rpr += '<w:i/>'
    rpr += '</w:rPr>'
    return f'<w:r>{rpr}<w:t xml:space="preserve">{esc(text)}</w:t></w:r>'


def make_para_runs(text, size_half=22):
    """Parse inline **bold** / *italic* and return XML runs."""
    runs = ''
    for m in INLINE_RE.finditer(text):
        bt = m.group(2)
        it = m.group(3)
        pt = m.group(4)
        if bt:
            runs += make_run(bt, bold=True,   size_half=size_half)
        elif it:
            runs += make_run(it, italic=True,  size_half=size_half)
        elif pt:
            runs += make_run(pt, size_half=size_half)
    return runs


def make_paragraph(text, style='Normal', size_half=22):
    ppr = f'<w:pPr><w:pStyle w:val="{style}"/><w:spacing w:after="120" w:line="280" w:lineRule="auto"/></w:pPr>'
    return f'<w:p>{ppr}{make_para_runs(text, size_half)}</w:p>'


def make_heading(text, level):
    style = f'Heading{level}'
    sizes = {1: 32, 2: 28, 3: 24, 4: 22}
    sh    = sizes.get(level, 22)
    ppr   = f'<w:pPr><w:pStyle w:val="{style}"/><w:spacing w:before="240" w:after="80"/></w:pPr>'
    runs  = make_para_runs(text, size_half=sh)
    # force bold
    runs  = re.sub(r'(<w:rPr>)', r'\1<w:b/>', runs)
    return f'<w:p>{ppr}{runs}</w:p>'


def make_table(rows):
    ncols = max(len(r) for r in rows)
    border_xml = ('<w:tblBorders>'
                  + ''.join(f'<w:{n} w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                             for n in ('top','left','bottom','right','insideH','insideV'))
                  + '</w:tblBorders>')
    tbl = f'<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/>{border_xml}</w:tblPr>'
    for ri, row in enumerate(rows):
        tbl += '<w:tr>'
        for ci in range(ncols):
            txt  = row[ci] if ci < len(row) else ''
            fill = 'D9D9D9' if ri == 0 else 'FFFFFF'
            shd  = f'<w:shd w:val="clear" w:color="auto" w:fill="{fill}"/>'
            tcpr = f'<w:tcPr>{shd}</w:tcPr>'
            runs = make_para_runs(txt, size_half=20)
            if ri == 0:
                runs = re.sub(r'(<w:rPr>)', r'\1<w:b/>', runs)
            tbl += f'<w:tc>{tcpr}<w:p><w:pPr><w:spacing w:after="60"/></w:pPr>{runs}</w:p></w:tc>'
        tbl += '</w:tr>'
    tbl += '</w:tbl>'
    return tbl


def build_docx_xml(lines):
    """Build the w:document body XML from lines."""
    body = ''
    i    = 0
    while i < len(lines):
        s = lines[i].strip()

        if not s:
            i += 1
            continue

        if re.match(r'^-{3,}$', s):
            body += '<w:p><w:pPr><w:pBdr><w:bottom w:val="single" w:sz="6" w:space="1" w:color="AAAAAA"/></w:pBdr><w:spacing w:before="120" w:after="120"/></w:pPr></w:p>'
            i += 1
            continue

        hm = re.match(r'^(#{1,4})\s+(.*)', s)
        if hm:
            body += make_heading(hm.group(2).strip(), len(hm.group(1)))
            i += 1
            continue

        if s.startswith('|'):
            rows, i = parse_table(lines, i)
            if rows:
                body += make_table(rows)
                body += '<w:p><w:pPr><w:spacing w:after="80"/></w:pPr></w:p>'
            continue

        body += make_paragraph(s)
        i += 1

    return body


CONTENT_TYPES = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml"  ContentType="application/xml"/>
  <Override PartName="/word/document.xml"
    ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/styles.xml"
    ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
  <Override PartName="/word/settings.xml"
    ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"/>
</Types>'''

RELS = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1"
    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
    Target="word/document.xml"/>
</Relationships>'''

WORD_RELS = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1"
    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles"
    Target="styles.xml"/>
  <Relationship Id="rId2"
    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/settings"
    Target="settings.xml"/>
</Relationships>'''

SETTINGS = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:settings xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:defaultTabStop w:val="720"/>
</w:settings>'''

STYLES = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
          xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">
  <w:docDefaults>
    <w:rPrDefault><w:rPr>
      <w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Calibri"/>
      <w:sz w:val="22"/><w:szCs w:val="22"/>
      <w:lang w:val="de-DE"/>
    </w:rPr></w:rPrDefault>
  </w:docDefaults>
  <w:style w:type="paragraph" w:styleId="Normal" w:default="1">
    <w:name w:val="Normal"/>
    <w:pPr><w:spacing w:after="120" w:line="280" w:lineRule="auto"/></w:pPr>
    <w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="22"/></w:rPr>
  </w:style>
  <w:style w:type="paragraph" w:styleId="Heading1">
    <w:name w:val="heading 1"/>
    <w:basedOn w:val="Normal"/>
    <w:pPr><w:outlineLvl w:val="0"/>
      <w:spacing w:before="280" w:after="80"/>
    </w:pPr>
    <w:rPr><w:b/><w:sz w:val="32"/><w:szCs w:val="32"/></w:rPr>
  </w:style>
  <w:style w:type="paragraph" w:styleId="Heading2">
    <w:name w:val="heading 2"/>
    <w:basedOn w:val="Normal"/>
    <w:pPr><w:outlineLvl w:val="1"/>
      <w:spacing w:before="240" w:after="80"/>
    </w:pPr>
    <w:rPr><w:b/><w:sz w:val="28"/><w:szCs w:val="28"/></w:rPr>
  </w:style>
  <w:style w:type="paragraph" w:styleId="Heading3">
    <w:name w:val="heading 3"/>
    <w:basedOn w:val="Normal"/>
    <w:pPr><w:outlineLvl w:val="2"/>
      <w:spacing w:before="200" w:after="60"/>
    </w:pPr>
    <w:rPr><w:b/><w:sz w:val="24"/><w:szCs w:val="24"/></w:rPr>
  </w:style>
  <w:style w:type="paragraph" w:styleId="Heading4">
    <w:name w:val="heading 4"/>
    <w:basedOn w:val="Normal"/>
    <w:pPr><w:outlineLvl w:val="3"/>
      <w:spacing w:before="160" w:after="40"/>
    </w:pPr>
    <w:rPr><w:b/><w:sz w:val="22"/><w:szCs w:val="22"/></w:rPr>
  </w:style>
  <w:style w:type="table" w:styleId="TableGrid">
    <w:name w:val="Table Grid"/>
    <w:tblPr>
      <w:tblBorders>
        <w:top    w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:left   w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:right  w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>
      </w:tblBorders>
    </w:tblPr>
    <w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="20"/></w:rPr>
  </w:style>
</w:styles>'''


def build_fallback(md_path, docx_path):
    """Build a valid .docx using raw XML + zipfile (no python-docx needed)."""
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().split('\n')

    body_xml = build_docx_xml(lines)

    # Page size A4, margins in twips (1cm = 567 twips)
    # top=2cm=1134, bottom=2cm=1134, left=2.5cm=1417, right=2.5cm=1417
    doc_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas"
  xmlns:mo="http://schemas.microsoft.com/office/mac/office/2008/main"
  xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
  xmlns:mv="urn:schemas-microsoft-com:mac:vml"
  xmlns:o="urn:schemas-microsoft-com:office:office"
  xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
  xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"
  xmlns:v="urn:schemas-microsoft-com:vml"
  xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing"
  xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
  xmlns:w10="urn:schemas-microsoft-com:office:word"
  xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
  xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"
  xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"
  xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk"
  xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml"
  xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
  mc:Ignorable="w14 wp14">
  <w:body>
{body_xml}
    <w:sectPr>
      <w:pgSz w:w="11906" w:h="16838"/>
      <w:pgMar w:top="1134" w:right="1417" w:bottom="1134" w:left="1417"
               w:header="709" w:footer="709" w:gutter="0"/>
    </w:sectPr>
  </w:body>
</w:document>'''

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        zf.writestr('[Content_Types].xml', CONTENT_TYPES)
        zf.writestr('_rels/.rels',          RELS)
        zf.writestr('word/_rels/document.xml.rels', WORD_RELS)
        zf.writestr('word/document.xml',    doc_xml)
        zf.writestr('word/styles.xml',      STYLES)
        zf.writestr('word/settings.xml',    SETTINGS)

    with open(docx_path, 'wb') as f:
        f.write(buf.getvalue())


# ═══════════════════════════════════════════════════════════════════════════
# ENTRY POINT
# ═══════════════════════════════════════════════════════════════════════════

if __name__ == '__main__':
    print(f'Converting:\n  {MD_PATH}\n  → {DOCX_PATH}')
    if HAS_DOCX:
        print('Using python-docx ...')
        build_with_docx(MD_PATH, DOCX_PATH)
    else:
        print('python-docx not found — using raw XML fallback ...')
        build_fallback(MD_PATH, DOCX_PATH)

    size_kb = os.path.getsize(DOCX_PATH) // 1024
    print(f'[OK] Written: {DOCX_PATH}  ({size_kb} KB)')
    print('FERTIG')
